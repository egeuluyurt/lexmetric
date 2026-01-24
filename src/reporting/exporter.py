from docxtpl import DocxTemplate
from docx import Document
from io import BytesIO
import pandas as pd
from datetime import datetime
import os
import logging
from src.audit_engine.config import REGULATORY_DATA_2026

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class ReportExporter:
    """
    Generates professional Audit Defense Reports using docxtpl.
    """
    
    TEMPLATE_PATH = "audit_report_template.docx"

    @classmethod
    def _ensure_template_exists(cls):
        """Creates a basic Jinja2-ready .docx template if missing."""
        if not os.path.exists(cls.TEMPLATE_PATH):
            doc = Document()
            # Dynamic Header based on 2026 PA Launch
            doc.add_heading('Medicaid Defense Audit - Commonwealth of {{ state_code }}', 0)
            doc.add_paragraph('Generated on: {{ generation_date }}')
            
            doc.add_heading('1. Executive Summary', level=1)
            doc.add_paragraph('Total Transactions Reviewed: {{ total_txns }}')
            doc.add_paragraph('Look-back Period: {{ start_date }} to {{ end_date }}')
            doc.add_paragraph('State Jurisdiction: {{ state_code }}')
            
            doc.add_heading('2. Regulatory Truth (2026 Rules)', level=1)
            doc.add_paragraph('Daily Penalty Rate: ${{ daily_rate }}')
            doc.add_paragraph('Asset Limit: ${{ asset_limit }}')
            doc.add_paragraph('Note: {{ regulatory_note }}')

            doc.add_heading('3. Forensic Risk Assessment', level=1)
            doc.add_paragraph('Total Unallowable Transfers: ${{ total_unallowable }}')
            doc.add_paragraph('Calculated Penalty Period: {{ penalty_days }} days ({{ penalty_months }} months)')
            
            doc.add_heading('4. Flagged Transactions Detail', level=1)
            doc.add_paragraph('The following transactions were identified as High Risk:')
            
            # Table placeholder for 'transactions'
            # In a real template we'd use complex jinja tags inside the word doc table
            # For this generated one, we will just use a loop placeholder in text form for simplicity
            doc.add_paragraph('{% for txn in transactions %}')
            doc.add_paragraph('Date: {{ txn.date }} | Amount: ${{ txn.amount }} | Desc: {{ txn.desc }}')
            doc.add_paragraph('Reasoning: {{ txn.reasoning }}')
            doc.add_paragraph('--------------------------------------------------')
            doc.add_paragraph('{% endfor %}')
            
            doc.save(cls.TEMPLATE_PATH)

    @classmethod
    def generate_report(cls, 
                        ledger: pd.DataFrame, 
                        audit_context: dict,
                        penalty_results: tuple) -> BytesIO:
        """
        Renders the report and returns a BytesIO stream.
        """
        cls._ensure_template_exists()
        
        try:
            doc = DocxTemplate(cls.TEMPLATE_PATH)
            
            # Prepare Data Context
            state = audit_context.get('state_rules', 'PA')
            reg_data = REGULATORY_DATA_2026.get(state, {})
            
            # Filter High Risk Items
            high_risk = ledger[ledger['Risk_Level'] >= 7].copy()
            
            # Format Transactions for Template
            txns_list = []
            total_unallowable = 0.0
            
            for _, row in high_risk.iterrows():
                amt = float(row.get('Withdrawal', 0)) or float(row.get('Amount', 0))
                total_unallowable += amt
                txns_list.append({
                    "date": str(row['Date']),
                    "amount": f"{amt:,.2f}",
                    "desc": row['Description'],
                    "reasoning": row.get('Forensic_Reasoning', 'N/A')
                })

            penalty_days, penalty_months = penalty_results

            context = {
                "generation_date": datetime.now().strftime("%Y-%m-%d"),
                "total_txns": len(ledger),
                "start_date": audit_context.get('lookback_start', 'N/A'),
                "end_date": audit_context.get('lookback_end', 'N/A'),
                "state_code": state,
                "daily_rate": reg_data.get('daily_penalty_divisor', 'N/A'),
                "asset_limit": reg_data.get('asset_limit', 'N/A'),
                "regulatory_note": f"Calculated based on 2026 {state} State Rate of ${reg_data.get('daily_penalty_divisor')}/day",
                "total_unallowable": f"{total_unallowable:,.2f}",
                "penalty_days": penalty_days,
                "penalty_months": penalty_months,
                "transactions": txns_list
            }
            
            doc.render(context)
            
            # Save to BytesIO
            output_stream = BytesIO()
            doc.save(output_stream)
            output_stream.seek(0)
            return output_stream

        except Exception as e:
            logger.error(f"Report Generation Failed: {e}")
            raise e
